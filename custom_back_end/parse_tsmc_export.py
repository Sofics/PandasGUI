import datetime
import datetime as dt
import re
from pathlib import Path
from pprint import pprint

from openpyxl.reader.excel import load_workbook

from teggydb import WaferVolume


class SimpleWaferVolume:
    """Utility class to save data and simply get rid of duplicates /w a set."""

    def __init__(self, ip_name: str, ip_version: str, tapeouts: int, wafers: int, export_date: dt.date = None):
        self.ip_name = ip_name  # equivalent to "product" in delivered cells
        self.ip_version = ip_version  # equivalent to "tag" in delivered cells
        self.tapeouts = tapeouts
        self.wafers = wafers
        self.export_date = export_date

    def __hash__(self):
        return hash(self.ip_name + self.ip_version + str(self.export_date))

    def __eq__(self, other):
        if isinstance(self, SimpleWaferVolume) and isinstance(other, SimpleWaferVolume):
            return self.__hash__() == other.__hash__()
        else:
            return False


def parse_tsmc_export_to_wafer_volumes(tsmc_export_path: Path):
    """Function to extract data from a TSMC wafer report and add it to the Teggy DB (WaferVolume table)."""

    workbook = load_workbook(tsmc_export_path)
    try:
        sheet = workbook["Summary-V991"]
    except KeyError:
        sheet = workbook.worksheets[1]

    # Finding relevant columns & the row_nr where the data starts:
    ip_name_coordinates = None
    ip_version_coordinates = None
    volume_production_coordinates = None
    row_nr = 0
    for row in sheet.rows:
        row_nr += 1
        if None not in (ip_name_coordinates, ip_version_coordinates, volume_production_coordinates):
            break
        col_nr = 0
        for cell in row:
            col_nr += 1
            if cell.value is not None:
                if cell.value.strip().lower().replace("\n", " ").startswith("ip name"):
                    ip_name_coordinates = (row_nr, col_nr)
                elif cell.value.strip().lower().replace("\n", " ").startswith("ip ver"):
                    ip_version_coordinates = (row_nr, col_nr)
                elif cell.value.strip().lower().replace("\n", " ").startswith("volume production"):
                    volume_production_coordinates = (row_nr, col_nr)
    # print(ip_name_coordinates, ip_version_coordinates, volume_production_coordinates)

    # creating unique SimpleWaverVolumes:
    simple_wafervolumes = set()
    for row_nr in range(ip_name_coordinates[0] + 1, sheet.max_row + 1):
        ip_name = sheet.cell(row_nr, ip_name_coordinates[1]).value
        ip_version = sheet.cell(row_nr, ip_version_coordinates[1]).value
        volume_production_str = sheet.cell(row_nr, volume_production_coordinates[1]).value
        splits = volume_production_str.replace(" ", "").split("/")
        tapeouts = int(splits[0])
        wafers = int(splits[1])

        simple_wafervolume = SimpleWaferVolume(ip_name, ip_version, tapeouts, wafers)
        simple_wafervolumes.add(simple_wafervolume)
    # print(len(simple_wafervolumes))

    # determine export date:
    file_name = tsmc_export_path.name
    match = re.search(r"_(\d{8})", file_name)
    if match:
        # fetch export date from filename
        export_date = dt.datetime.strptime(match.group(1), "%Y%m%d").date()
    else:
        # fetch export date from first tab, field B3
        workbook = load_workbook(tsmc_export_path)
        sheet = workbook.worksheets[0]
        date_str = sheet["B3"].value.split(" ")[0]
        export_date = dt.datetime.strptime(date_str, "%Y/%m/%d").date()

    # create the WaferVolumes (& automatically add them to the database)
    for swv in simple_wafervolumes:
        WaferVolume(export_date, swv.ip_name, swv.ip_version, swv.tapeouts, swv.wafers)


def one_time_tsmc_wafer_history_import(tsmc_wafer_followup_path: Path):
    """This function imports historical TSMC wafer data from Bart K's WaferFollowUp.xlsx to the Teggy DB."""

    workbook = load_workbook(tsmc_wafer_followup_path)
    try:
        sheet = workbook["Sheet1"]
    except KeyError:
        sheet = workbook.worksheets[0]

    # NOTE: all volumeproduction columns we'd like to import have their title transformed to YYYYMMDD
    column_nr_export_date_dict = {}  # col_nr -> export date
    for row in sheet.rows:  # only do the first row
        col_nr = 0
        for cell in row:
            col_nr += 1
            if type(cell.value) == datetime.datetime:
                # TODO format all columns before oct 2022 to be dates and process them (!!! with adjusted wafer count!!)
                pass
                # # automatically found as date by openpyxl...
                # export_date = cell.value.date()
                # column_nr_export_date_dict.update({col_nr: export_date})
            elif cell.value is not None:
                match = re.search(r"(\d{8})", cell.value)
                if match:
                    # TODO Note: for now we only use columns I manually formatted to YYYYMMDD
                    #  These are the columns from oct 2022 onwards, July 2022 -> Oct 2022 a wafer adjustment happened.
                    export_date = dt.datetime.strptime(match.group(1), "%Y%m%d").date()
                    column_nr_export_date_dict.update({col_nr: export_date})
        break  # only do the first row
    # pprint(column_nr_export_date_dict)

    # creating unique SimpleWaverVolumes:
    simple_wafervolumes = set()
    row_nr = 1  # skip the first row
    for row in sheet.rows:
        row_nr += 1
        ip_name = str(sheet.cell(row_nr, 1).value).strip()
        ip_version = str(sheet.cell(row_nr, 2).value).strip()
        for column_nr, export_date in column_nr_export_date_dict.items():
            volume_production_str = sheet.cell(row_nr, column_nr).value
            if volume_production_str is not None:
                splits = volume_production_str.replace(" ", "").split("/")
                if len(splits) == 2:
                    tapeouts = int(splits[0])
                    wafers = int(splits[1])
                    swv = SimpleWaferVolume(ip_name, ip_version, tapeouts, wafers, export_date)
                    simple_wafervolumes.add(swv)
    # print(len(simple_wafervolumes))

    # create the WaferVolumes (& automatically add them to the database)
    for swv in simple_wafervolumes:
        WaferVolume(swv.export_date, swv.ip_name, swv.ip_version, swv.tapeouts, swv.wafers)


if __name__ == "__main__":
    pass
    # parse_tsmc_export_to_wafer_volumes(Path(r"C:\repos\PandasGUI\custom_back_end\Export_20241001193141.xlsx"))
    one_time_tsmc_wafer_history_import(Path(r"C:\repos\PandasGUI\custom_back_end\2024_10_WaferFollowUp.xlsx"))
